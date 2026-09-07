import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bảng CDC StarRocks"
    ws.views.sheetView[0].showGridLines = True

    # Title
    ws.merge_cells("A1:E1")
    title_cell = ws["A1"]
    title_cell.value = "TỔNG HỢP CÁC BẢNG DATABASE CDC TRÊN STARROCKS (kfm_scm)"
    title_cell.font = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    # Subtitle
    ws.merge_cells("A2:E2")
    sub_cell = ws["A2"]
    sub_cell.value = "Máy chủ: 103.147.122.103:9030 | Database: kfm_scm | Cập nhật: 04/09/2026 | Quy mô: 26 bảng CDC (~225+ triệu dòng)"
    sub_cell.font = Font(name="Arial", size=10, italic=True, color="333333")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    # Headers
    headers = ["STT", "Nhóm Nghiệp Vụ", "Tên Bảng CDC (StarRocks)", "Số Lượng Dòng", "Ý Nghĩa & Ứng Dụng Nghiệp Vụ"]
    ws.append([]) # Row 3 blank
    ws.append(headers) # Row 4
    ws.row_dimensions[4].height = 28

    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_border = Border(
        left=Side(style='thin', color='B0C4DE'),
        right=Side(style='thin', color='B0C4DE'),
        top=Side(style='medium', color='1F4E78'),
        bottom=Side(style='medium', color='1F4E78')
    )

    for col_idx in range(1, 6):
        cell = ws.cell(row=4, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border

    # Data Rows
    data = [
        # Nhóm 1
        (1, "Quản lý Kho & Tồn kho", "__cdc_kfm_kf_inventories_kf_inventory_transaction_stock_summaries", 155820973, "Tổng hợp số lượng xuất, nhập, tồn kho định kỳ theo từng SKU toàn hệ thống."),
        (2, "Quản lý Kho & Tồn kho", "__cdc_kfm_kf_inventories_kf_inventory_transaction_stockcard", 50584173, "Sổ thẻ kho chi tiết từng giao dịch làm biến động tồn kho (thời gian, mã phiếu, +/- số lượng)."),
        (3, "Quản lý Kho & Tồn kho", "__cdc_kfm_kf_inventories_kf_transfer_items", 315096, "Danh mục các mặt hàng được điều chuyển giữa DC và các Cửa hàng/Siêu thị."),
        (4, "Quản lý Kho & Tồn kho", "__cdc_kfm_kf_inventories_kf_transfer_items___container_lines", 1259418, "⭐ CỐT LÕI ĐỐI SOÁT THÙNG RỔ: Chi tiết từng mã thùng/tote/container, thời gian quét nhận, SL chuyển vs SL nhận."),
        (5, "Quản lý Kho & Tồn kho", "__cdc_kfm_kf_inventories_kf_transfer_items___return_transfers", 67427, "Dữ liệu các đợt hàng hoặc thùng rổ chuyển hoàn ngược từ Siêu thị về DC."),
        (6, "Quản lý Kho & Tồn kho", "__cdc_kfm_kf_inventories_kf_transfer_items___trips", 74914, "Bảng liên kết trung gian giữa mặt hàng luân chuyển và mã chuyến xe vận chuyển."),
        (7, "Quản lý Kho & Tồn kho", "__cdc_kfm_kf_inventories_kf_claim_stock_summaries", 2, "Dữ liệu tổng hợp các phiếu khiếu nại bồi hoàn chênh lệch tồn kho."),
        (8, "Quản lý Kho & Tồn kho", "__cdc_kfm_kf_inventories_kf_claim_transaction_details", 2, "Chi tiết các giao dịch bồi hoàn/xử lý khiếu nại chênh lệch tồn."),
        # Nhóm 2
        (9, "Vận Tải & Giao Nhận", "__cdc_kfm_kf_inventories_kf_trips", 16009, "Danh sách các chuyến xe vận tải điều phối hàng ngày từ DC đến các điểm giao."),
        (10, "Vận Tải & Giao Nhận", "__cdc_kfm_kf_inventories_kf_trip_locations", 57215, "Các điểm dừng giao hàng theo chuyến (Cửa hàng, Hub trung chuyển, DC)."),
        (11, "Vận Tải & Giao Nhận", "__cdc_kfm_kf_inventories_kf_trip_locations___tl_received_pts", 823305, "⭐ CỐT LÕI ĐỐI SOÁT SLDT: Danh sách các Phiếu Điều Chuyển (PT) được bàn giao thực tế tại từng ST."),
        (12, "Vận Tải & Giao Nhận", "__cdc_kfm_kf_inventories_kf_trips_locations_items", 387203, "Chi tiết số lượng từng mã hàng được giao tại từng điểm dừng của xe."),
        (13, "Vận Tải & Giao Nhận", "__cdc_kfm_597a32ea_597a32ea_L1___tl_driver_received_pts", 95223, "Xác nhận ký nhận phiếu điều chuyển của tài xế phụ trách chuyến xe."),
        (14, "Vận Tải & Giao Nhận", "__cdc_kfm_597a32ea_597a32ea_L1___tl_temperature_images", 193221, "Hình ảnh và nhật ký nhiệt độ thùng bảo quản lạnh trên xe vận chuyển."),
        # Nhóm 3
        (15, "Chứng Từ Điều Chuyển", "__cdc_kfm_kf_transfer_tickets_kf_transfer_tickets", 51504, "Phiếu điều chuyển gốc từ WMS/ERP (mã PT, mã seal, mã container, trạng thái)."),
        (16, "Chứng Từ Điều Chuyển", "__cdc_kfm_kf_transfer_tickets_kf_transfer_ticket_lines", 162959, "Chi tiết từng SKU, quy cách đóng gói và số lượng xuất trên phiếu điều chuyển."),
        # Nhóm 4
        (17, "Kiểm Soát Chất Lượng (QC/HRW)", "__cdc_kfm_kf_inventories_kf_hrw_quality_tickets", 21896, "Phiếu thẩm định kiểm tra chất lượng ngành hàng tươi sống/rau củ quả (HRW)."),
        (18, "Kiểm Soát Chất Lượng (QC/HRW)", "__cdc_kfm_kf_inventories_kf_hrw_quality_tickets___items", 101657, "Chi tiết tình trạng kiểm tra theo từng mã hàng (tươi, dập nát, hư hỏng)."),
        (19, "Kiểm Soát Chất Lượng (QC/HRW)", "__cdc_kfm_bede5be7_de5b9963_L2___criterias", 1030745, "Danh mục bộ tiêu chuẩn kỹ thuật đánh giá quy cách, độ tươi và lỗi hao hụt."),
        (20, "Kiểm Soát Chất Lượng (QC/HRW)", "__cdc_kfm_bede5be7_de5b9963_L2___evaluations", 41640, "Ghi nhận số lượng lỗi nặng (heavy defect), lỗi nhẹ (light defect) khi nhập kho."),
        # Nhóm 5
        (21, "Line Items & Đơn Hàng", "__cdc_kfm_ec9d24ab_1a050070_L2___product_lines", 8283134, "Phân rã dữ liệu đơn hàng theo từng dòng sản phẩm."),
        (22, "Line Items & Đơn Hàng", "__cdc_kfm_ec9d24ab_33bc7bbc_L3___line_items", 8281829, "Dữ liệu mặt hàng chi tiết nhất của các đơn đặt hàng luân chuyển."),
        (23, "Line Items & Đơn Hàng", "__cdc_kfm_ec9d24ab_e99d544c_L4___origin_transfer_codes", 188040, "Truy vết nguồn gốc đơn hàng từ mã phiếu chuyển kho khởi tạo ban đầu."),
        (24, "Line Items & Đơn Hàng", "__cdc_kfm_ec9d24ab_c49f2956_L2___trip_locations", 74914, "Vị trí giao nhận gắn với từng phân dòng sản phẩm."),
        (25, "Line Items & Đơn Hàng", "__cdc_kfm_c6c82317_c6c82317_L1___created_quantity_histories", 22311, "Ghi vết các lần điều chỉnh số lượng tạo đơn ban đầu của người lập phiếu."),
        # Nhóm 6
        (26, "Quản Trị Metadata CDC", "__cdc_schema_changes", 12, "Tự động ghi vết các trường dữ liệu được cập nhật hoặc đổi kiểu từ nguồn CDC sang StarRocks.")
    ]

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    star_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    start_row = 5
    for i, row in enumerate(data):
        current_row = start_row + i
        ws.append(list(row))
        ws.row_dimensions[current_row].height = 24
        
        is_star = "⭐" in row[4]
        is_even = (i % 2 == 1)

        for col_idx in range(1, 6):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = Font(name="Arial", size=10)
            cell.border = thin_border
            
            if is_star:
                cell.fill = star_fill
                if col_idx in [3, 5]:
                    cell.font = Font(name="Arial", size=10, bold=True, color="C00000")
            elif is_even:
                cell.fill = zebra_fill

            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 3:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 4:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0'
            elif col_idx == 5:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Column widths
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 65
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 85

    # Total Row
    total_row = start_row + len(data)
    ws.merge_cells(f"A{total_row}:C{total_row}")
    t_cell = ws[f"A{total_row}"]
    t_cell.value = "TỔNG CỘNG TẤT CẢ BẢN GHI CDC TRÊN STARROCKS"
    t_cell.font = Font(name="Arial", size=11, bold=True, color="1F4E78")
    t_cell.alignment = Alignment(horizontal="right", vertical="center")
    
    sum_cell = ws[f"D{total_row}"]
    sum_cell.value = f"=SUM(D{start_row}:D{total_row-1})"
    sum_cell.font = Font(name="Arial", size=11, bold=True, color="1F4E78")
    sum_cell.number_format = '#,##0'
    sum_cell.alignment = Alignment(horizontal="right", vertical="center")

    for col in range(1, 6):
        c = ws.cell(row=total_row, column=col)
        c.border = Border(top=Side(style='medium', color='1F4E78'), bottom=Side(style='double', color='1F4E78'))
        c.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    ws.row_dimensions[total_row].height = 28

    out_file = "h:/My Drive/Lynguyen/TONG_HOP_DATABASE_CDC_STARROCKS.xlsx"
    wb.save(out_file)
    print(f"Exported Excel successfully to {out_file}")

if __name__ == "__main__":
    generate_excel()
